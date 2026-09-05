#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o .docx da Analise Diagnostica do Ciclo 1 (POPs 5, 6 e 7).

Fonte: dados/outputs/dre-fluxo-receita.json, dre-diagnostico-travas.json e
dre-maturidade-digital.json. O .docx e gerado, nunca editado a mao: quando o dado
mudar, regenere.

Nota: identificadores e comentarios sem acento de proposito. O texto que vai para o
cliente, dentro das strings, e acentuado.
"""

import json, os, re, pathlib
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from importlib.machinery import SourceFileLoader

AQUI = pathlib.Path(os.path.dirname(os.path.abspath(__file__)))
RAIZ = AQUI.parent.parent
_g = SourceFileLoader("gera_docx", str(AQUI / "gera-docx.py")).load_module()

setup, cabecalho, h2, h3, para = _g.setup, _g.cabecalho, _g.h2, _g.h3, _g.para
rule, panel, bullets, run = _g.rule, _g.panel, _g.bullets, _g.run
shade, cell_margins, borders = _g.shade, _g.cell_margins, _g.borders
INK, MUTED, FAINT, OK, WARN, CRIT = _g.INK, _g.MUTED, _g.FAINT, _g.OK, _g.WARN, _g.CRIT
OXDARK, SUNK, BRANDWASH, OKWASH, MONO = _g.OXDARK, _g.SUNK, _g.BRANDWASH, _g.OKWASH, _g.MONO


def carrega(nome):
    return json.loads((RAIZ / "dados" / "outputs" / f"{nome}.json").read_text())


def limpo(s, n=None):
    """Achata quebras de linha e corta marcacao que nao sobrevive ao Word."""
    s = re.sub(r"`([^`]*)`", r"\1", str(s))
    s = re.sub(r"\*\*([^*]*)\*\*", r"\1", s)
    s = re.sub(r"\s*\n\s*", " ", s).strip()
    return (s[: n - 1] + "…") if n and len(s) > n else s


def tabela(doc, headers, rows, larguras):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    borders(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        shade(c, SUNK)
        cell_margins(c)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run(p, h.upper(), size=8, bold=True, color=FAINT)
    for r in rows:
        cs = t.add_row().cells
        for i, val in enumerate(r):
            cs[i].text = ""
            cell_margins(cs[i])
            txt, kind = (val if isinstance(val, tuple) else (val, "n"))
            p = cs[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.14
            if kind == "mono":
                run(p, txt, size=8.5, color=MUTED, font=MONO)
            elif kind == "b":
                run(p, txt, size=9, bold=True, color=INK)
            elif kind == "null":
                run(p, txt, size=8.5, color=FAINT, font=MONO)
            else:
                run(p, txt, size=8.8, color=INK)
    for row in t.rows:
        for i, w in enumerate(larguras):
            row.cells[i].width = Cm(w)
    para(doc, "", space_after=2)


def doc_diagnostico(path):
    f, t, m = carrega("dre-fluxo-receita"), carrega("dre-diagnostico-travas"), carrega("dre-maturidade-digital")
    doc = Document()
    setup(doc)
    cabecalho(doc, "Análise Diagnóstica",
              "Fluxo de receita, as 8 travas e a maturidade digital, Ciclo 1",
              "04/09/2026  ·  leitura parcial, anterior ao desbloqueio de acesso de 08/09")

    panel(doc, "Como ler este documento",
          ["Produzido com a evidência disponível em 04/09, antes da atualização de acessos "
           "comprometida pela OLX para 08/09. É leitura parcial e datada.",
           "Onde um campo aparece como “sem dado”, ele está assim por falta de dado e não por "
           "omissão: a regra do método é que campo sem fonte fica vazio e vira pendência, nunca "
           "um valor aproximado.",
           "Marcação de origem: [D] indica número declarado pelo Grupo OLX e ainda não apurado "
           "pela V4; [E] indica estimativa da V4. Número sem marca é apurado na ferramenta."])

    # ---------------------------------------------------------- POP 5
    h2(doc, "01", "Mapeamento do Fluxo de Receita")
    para(doc, limpo(f["summary_headline"]), size=11, bold=True, space_after=6)
    para(doc, limpo(f["summary"]), size=10, color=MUTED, space_after=8)

    h3(doc, "Em números")
    tabela(doc, ["Indicador", "O que significa"],
           [[(limpo(h["valor"]), "b"), limpo(h["subtext"], 260)] for h in f["summary_highlights"]],
           [4.0, 12.8])

    h3(doc, "Os achados")
    bullets(doc, [limpo(k, 700) for k in f["summary_key_findings"]])

    h3(doc, "As dez etapas")
    para(doc, "Origem: apurado = a V4 mediu · declarado [D] = o cliente afirmou e ninguém conferiu · "
              "inferido = leitura da V4 · ausente = não há dado.",
         size=9, color=MUTED, space_after=4)
    tabela(doc, ["#", "Etapa", "Origem", "Volume", "Perda absoluta", "Travas"],
           [[(str(e["ordem"]), "mono"), (limpo(e["nome"]), "b"), e["natureza"],
             (limpo(e["volume"] or "sem dado"), "null" if not e["volume"] else "n"),
             (limpo(e["perda_absoluta"] or "sem dado"), "null" if not e["perda_absoluta"] else "n"),
             ", ".join(e["travas_associadas"]) or "-"] for e in f["etapas"]],
           [1.0, 4.6, 1.9, 3.0, 2.4, 3.9])

    h3(doc, "Ranking de perda, indeterminável")
    if f["ranking_perda"]:
        panel(doc, "O ranking que o POP exige não pode ser produzido",
              [limpo(f["ranking_perda"][0]["impossivel_ordenar"], 1100)], fill=BRANDWASH)

    h3(doc, "Reconciliação de receita")
    para(doc, limpo(f["reconciliacao"]["explicacao"], 2200), size=9.5, space_after=6)

    h3(doc, "Ordem de diagnóstico das travas")
    para(doc, "Regra de Goldratt: de baixo para cima no funil. Resolver Exposição com Retenção "
              "quebrada só aumenta o custo do desperdício.", size=9, color=MUTED, space_after=4)
    tabela(doc, ["#", "Trava", "Por quê"],
           [[(str(x["ordem"]), "mono"), (x["trava"], "b"), limpo(x["justificativa"], 420)]
            for x in f["travas_prioritarias"]],
           [1.0, 3.2, 12.6])

    h3(doc, "O que falta para fechar")
    bullets(doc, [limpo(x, 450) for x in f["o_que_falta_para_fechar"]])

    doc.add_page_break()

    # ---------------------------------------------------------- POP 6
    h2(doc, "02", "Diagnóstico das 8 Travas de Receita")
    para(doc, "Cada trava tem cinco dimensões pontuadas de 0 a 5. A regra do método limita a nota a 3 "
              "quando não existe evidência formal, por isso o máximo praticável hoje é 3 por dimensão, "
              "e não 5. Ler uma soma bruta contra a faixa de 0 a 25 do playbook seria erro de escala.",
         size=10, color=MUTED, space_after=8)

    tabela(doc, ["Trava", "Dimensões com nota", "Soma", "Confiabilidade", "Escala"],
           [[(x["nome"], "b"),
             f"{sum(1 for d in x['dimensoes'] if d.get('nota') is not None)} de 5",
             (str(x["score_total"]) if x["score_total"] is not None else "sem dado",
              "n" if x["score_total"] is not None else "null"),
             x["confiabilidade"], limpo(x["score_parcial"], 230)] for x in t["travas"]],
           [3.0, 2.6, 1.6, 2.4, 7.2])

    h3(doc, "Gate de consolidação causal")
    panel(doc, "A consolidação causal não pode rodar hoje",
          [limpo(t["gate_consolidacao"], 2400)], fill=BRANDWASH)

    h3(doc, "Candidata a restrição")
    para(doc, limpo(t["candidata_a_restricao"], 2600), size=9.5, space_after=6)
    para(doc, "Por que ainda não é conclusivo", size=9.5, bold=True, space_after=3)
    para(doc, limpo(t["por_que_nao_e_conclusiva"], 2000), size=9.5, space_after=6)

    h3(doc, "Alerta, o denominador da publicação")
    panel(doc, "Ambiguidade aberta que atravessou o modelo de meta",
          [limpo(t["alerta_denominador_publicacao"], 2600)], fill=BRANDWASH)

    h3(doc, "Trava a trava")
    for x in t["travas"]:
        para(doc, x["nome"], size=10.5, bold=True, color=OXDARK, space_before=8, space_after=2)
        para(doc, limpo(x["score_parcial"], 240), size=8.6, color=MUTED, font=MONO, space_after=4)
        tabela(doc, ["Dimensão", "Nota", "Camada", "Evidência"],
               [[limpo(d["dimensao"], 60),
                 (str(d["nota"]) if d.get("nota") is not None else "-",
                  "n" if d.get("nota") is not None else "null"),
                 d["natureza"], limpo(d["evidencia"], 640)] for d in x["dimensoes"]],
               [3.2, 1.2, 2.0, 10.4])
        if x.get("politica_implicita"):
            para(doc, "Política implícita: " + limpo(x["politica_implicita"], 700),
                 size=9, color=INK, space_after=3)
        para(doc, "O que falta: " + limpo(x["o_que_falta"], 700), size=9, color=MUTED, space_after=4)

    doc.add_page_break()

    # ---------------------------------------------------------- POP 7
    h2(doc, "03", "Maturidade Digital")
    para(doc, "O POP é explícito: avaliação por evidências objetivas e práticas reais, não por discurso. "
              "Um grupo deste porte parece maduro no discurso; o que segue pontua pelo que a ferramenta "
              "mostra.", size=10, color=MUTED, space_after=8)

    tabela(doc, ["Dimensão", "Nível"],
           [[(d["dimensao"], "b"),
             f"{d.get('nivel')}, {limpo(d['nome_do_nivel'])}"] for d in m["dimensoes"]],
           [6.0, 10.8])

    h3(doc, "Leitura sistêmica")
    para(doc, limpo(m["leitura_sistemica"], 2600), size=9.5, space_after=6)
    if m.get("paradoxo"):
        panel(doc, "O paradoxo", [limpo(m["paradoxo"], 1600)], fill=BRANDWASH)

    for d in m["dimensoes"]:
        h3(doc, f"{d['dimensao']}, nível {d.get('nivel')}")
        para(doc, "O que sustenta", size=9.5, bold=True, color=OK, space_after=3)
        bullets(doc, [limpo(e, 480) for e in d["evidencias"]])
        para(doc, "O que contradiz", size=9.5, bold=True, color=CRIT, space_before=4, space_after=3)
        bullets(doc, [limpo(e, 480) for e in d.get("contra_evidencias", [])] or ["-"])
        para(doc, "Para subir de nível: " + limpo(d["o_que_falta"], 900),
             size=9, color=MUTED, space_before=4, space_after=6)

    # ---------------------------------------------------------- fecho
    doc.add_page_break()
    h2(doc, "04", "O que este documento não é")
    panel(doc, "Três limites que precisam ser ditos em voz alta",
          ["Não é a restrição do ciclo. Seis das oito travas não têm score fechado, e a régua do "
           "método bloqueia a consolidação causal enquanto isso for verdade. O que existe aqui é "
           "candidata, não restrição validada.",
           "Não fecha a regra de reconciliação. A receita derivada do funil não pode ser comparada "
           "com a declarada, porque nenhuma das dez etapas tem volume absoluto.",
           "Não é definitivo. Foi produzido antes do desbloqueio de acesso de 08/09. Quase tudo que "
           "está marcado como sem dado tem caminho de resolução conhecido, e a maior parte dele é "
           "extração do que já existe, não coleta nova."],
          fill=BRANDWASH)

    _g.rodape(doc, "Registro de máquina em dados/outputs/ e versão longa em 02-diagnostico/.")
    doc.save(path)
    return path


if __name__ == "__main__":
    dest = os.path.expanduser("~/Desktop")
    p = doc_diagnostico(os.path.join(dest, "V4 x Grupo OLX - Analise Diagnostica.docx"))
    print("%9d bytes  %s" % (os.path.getsize(p), p))
